from __future__ import annotations

import os
import sys
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(85, 85, 85)
BLUE = RGBColor(46, 116, 181)
NAVY = RGBColor(31, 77, 120)
LIGHT_GRAY = RGBColor(242, 244, 247)


def set_run_font(run, *, name="Calibri", size=11, color=BLACK, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_separate, fld_text, fld_end])


def configure_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, NAVY, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1


def add_header_footer(doc, prepared_on):
    for section in doc.sections:
        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        left = paragraph.add_run("Sandworth Homes Board Brief")
        set_run_font(left, size=9, color=GRAY, bold=True)

        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_p.paragraph_format.space_after = Pt(0)
        run = footer_p.add_run(f"Prepared {prepared_on} | Page ")
        set_run_font(run, size=9, color=GRAY)
        add_page_number(footer_p)


def add_title_block(doc, prepared_on):
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(0)
    kicker.paragraph_format.space_after = Pt(2)
    kicker_run = kicker.add_run("BOARD BRIEF")
    set_run_font(kicker_run, size=10.5, color=GRAY, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("Sandworth Homes Platform Overview")
    set_run_font(title_run, size=23, color=BLACK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle.add_run("Current operating capability, user value, and near-term roadmap as of August 28, 2026")
    set_run_font(subtitle_run, size=13.5, color=GRAY)

    metadata = [
        ("Prepared for", "Board members and executive stakeholders"),
        ("Prepared by", "Product and platform review"),
        ("Date", prepared_on),
        ("Platform position", "Live housing and property marketplace with Phase 1 workflow upgrades"),
        ("Purpose", "Explain what the application does today, what value it creates, and what features are coming next"),
    ]

    for label, value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.1
        label_run = paragraph.add_run(f"{label}: ")
        set_run_font(label_run, size=11, color=BLACK, bold=True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=11, color=BLACK)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    note.paragraph_format.space_after = Pt(10)
    note.paragraph_format.left_indent = Inches(0.12)
    note.paragraph_format.right_indent = Inches(0.12)
    note.paragraph_format.line_spacing = 1.1
    note_run = note.add_run(
        "This brief is written for non-technical decision-makers. It explains the platform in business terms while keeping the current product scope and roadmap accurate."
    )
    set_run_font(note_run, size=10.5, color=NAVY)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(text)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color=BLACK)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.167
        if paragraph.runs:
            run = paragraph.runs[0]
            run.text = item
        else:
            run = paragraph.add_run(item)
        set_run_font(run, size=11, color=BLACK)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.167
        if paragraph.runs:
            run = paragraph.runs[0]
            run.text = item
        else:
            run = paragraph.add_run(item)
        set_run_font(run, size=11, color=BLACK)


def add_divider(doc):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D9DEE7")
    border.append(bottom)
    p_pr.append(border)


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.allow_autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F4F7")
    cell._tc.get_or_add_tcPr().append(shading)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=NAVY, bold=True)
    doc.add_paragraph()


def build_document(output_path):
    prepared_on = "August 28, 2026"
    doc = Document()
    configure_page(doc.sections[0])
    configure_styles(doc)
    add_header_footer(doc, prepared_on)

    add_title_block(doc, prepared_on)
    add_divider(doc)

    add_heading(doc, "Executive Summary", 1)
    add_body(
        doc,
        "Sandworth Homes is now a live property marketplace and operating platform for rentals, property sales, commercial leasing, and post-move tenant management. At its current level, the application does more than publish listings: it supports the customer journey from discovery through application, payment, tenancy record management, and operational follow-up."
    )
    add_body(
        doc,
        "The latest Phase 1 upgrades moved the platform closer to a real transaction and operations product. Renters can self-book tours from listing pages, buyers can submit offers, tenants can raise maintenance requests, users can message the team inside the app, and administrators can manage these workflows from one operating console."
    )
    add_callout(
        doc,
        "Board takeaway: the product has moved from a brochure-style listing site into an operational property platform, but it still needs deeper automation, real payment rails, and lease execution tools to become fully end-to-end."
    )

    add_heading(doc, "What The Platform Does Today", 1)
    add_body(
        doc,
        "The current application serves three groups at once: property seekers, existing tenants, and the internal operations team. It provides a shared digital workflow around housing and property transactions rather than separate disconnected tools."
    )
    add_bullets(
        doc,
        [
            "Public property discovery across homes for sale, annual rentals, and commercial spaces.",
            "Search-friendly listing pages with location, pricing, features, gallery, and share-ready property metadata.",
            "Account creation and sign-in for customers who want to apply, book tours, message the team, or manage a tenancy.",
            "Digital application flow for rental and commercial prospects, with online payment handoff when an application is approved.",
            "An internal admin area for inventory, applications, tour operations, offers, maintenance, and customer conversations.",
        ],
    )

    add_heading(doc, "Current User Value By Audience", 1)

    add_heading(doc, "For renters and tenants", 2)
    add_bullets(
        doc,
        [
            "Browse rental listings by location and property type, then open detailed property pages.",
            "Book a viewing slot directly from a rental property page when tour availability has been released.",
            "Submit a rental application online and move into payment immediately when auto-screening rules approve the case.",
            "Track applications and tours from a personal dashboard instead of relying on manual follow-up alone.",
            "After move-in, use the tenancy area to review ledger history, post payments, report maintenance issues, and message the operations team.",
            "Save search preferences so high-interest markets and matching inventory stay visible on the user dashboard.",
        ],
    )

    add_heading(doc, "For buyers", 2)
    add_bullets(
        doc,
        [
            "Browse residential sale listings with pricing, gallery, and location context.",
            "Submit an offer directly from a sale property page, including amount, timing, and terms.",
            "Keep buyer negotiations visible in the same account dashboard used for other housing journeys.",
        ],
    )

    add_heading(doc, "For commercial lease prospects", 2)
    add_bullets(
        doc,
        [
            "View commercial spaces with unit information, lease term, and location details.",
            "Submit a lease request and message the leasing desk from the property page.",
        ],
    )

    add_heading(doc, "For internal operations and management", 2)
    add_bullets(
        doc,
        [
            "Manage listing inventory, media, pricing, status, and availability from the admin side.",
            "Review applications, including instantly approved cases versus cases that still need manual review.",
            "Publish individual tour slots or generate a full block of viewing times for a property.",
            "Confirm, complete, or cancel upcoming tour requests from one tour desk.",
            "Receive buyer offers in a dedicated queue and update each offer through review, counter, acceptance, or decline.",
            "Handle property and tenancy messages in one inbox instead of scattered email threads.",
            "Track and resolve maintenance tickets through a dedicated maintenance queue.",
        ],
    )

    add_heading(doc, "Phase 1 Capability Added Recently", 1)
    add_body(
        doc,
        "The current build now includes the first must-have operational features that make the application more useful in real housing transactions and day-to-day management."
    )
    add_numbered(
        doc,
        [
            "Self-serve tour scheduling: admins publish availability and renters select an open slot themselves.",
            "Tour request management: customers can cancel and rebook, while admins can confirm or complete appointments.",
            "Basic automated screening: clearly qualified rental applicants can pass into the payment step without waiting for full manual handling.",
            "Buyer offer workflow: sale listings now support a direct offer path instead of forcing all negotiations offline.",
            "In-app messaging: users and admins can keep a conversation tied to the relevant property or tenancy.",
            "Maintenance ticketing: live tenants can raise issues and the operations team can track progress and resolution.",
            "Saved searches with live match counts: customers can monitor target areas and housing criteria more efficiently.",
        ],
    )

    add_heading(doc, "What The Product Still Does Not Fully Solve Yet", 1)
    add_body(
        doc,
        "Although the platform is materially stronger than a static listing site, it is not yet a fully closed-loop housing transaction system. Some important steps still rely on simplified logic or manual operational follow-up."
    )
    add_bullets(
        doc,
        [
            "The payment experience is still a simulated card form rather than a production integration with Paystack or Flutterwave.",
            "Screening logic is rule-based and light. It does not yet perform real income verification, guarantor checks, credit-like checks, or fraud controls.",
            "Tour reminders, customer notifications, and automatic slot release rules are still limited.",
            "Lease generation and e-signature are not yet built into the live workflow.",
            "Buyer journeys still need mortgage support, pre-qualification, and a stronger post-offer transaction flow.",
            "There is no live escrow or payment-protection layer yet for rent or purchase transactions.",
        ],
    )

    add_heading(doc, "Coming Features And Product Direction", 1)
    add_body(
        doc,
        "The next roadmap should focus on converting the current operational shell into a true end-to-end property transaction platform. The features below are the most commercially meaningful next steps."
    )

    add_heading(doc, "Next priority set for renters and tenants", 2)
    add_bullets(
        doc,
        [
            "Automatic tour release, confirmations, reminders, and no-show handling.",
            "Stronger automated screening with real verification checks and exception routing.",
            "Production payment gateway integration with sandbox and live modes.",
            "Lease generation, digital signing, and downloadable executed records.",
            "Richer conversation history, agent assignment, and notification rules inside the in-app chat.",
        ],
    )

    add_heading(doc, "Next priority set for buyers", 2)
    add_bullets(
        doc,
        [
            "Offer workflow expansion with counter-offer handling, document requests, and accepted-offer milestones.",
            "Mortgage and affordability calculator inside the planning flow.",
            "Pre-qualification intake tied to buyer readiness and lender follow-up.",
        ],
    )

    add_heading(doc, "Cross-platform operating features still to add", 2)
    add_bullets(
        doc,
        [
            "Escrow or payment-protection flows so customers can transact with more confidence online.",
            "Landlord or owner participation tools, including slot requests and operational responses.",
            "Deeper maintenance workflow with assignment, vendor tracking, status notifications, and service-level reporting.",
            "More complete reporting for conversion, response time, occupancy pipeline, and transaction completion.",
        ],
    )

    add_heading(doc, "Why This Matters To The Business", 1)
    add_bullets(
        doc,
        [
            "It shortens the gap between customer interest and action by turning listing pages into workflow entry points.",
            "It centralizes operations that would otherwise be split across phone calls, messaging apps, spreadsheets, and manual reminders.",
            "It improves visibility for management because tours, applications, offers, maintenance, and conversations become measurable platform activity.",
            "It creates a stronger foundation for revenue-generating features such as real payments, signed leases, premium property operations, and transaction support services.",
        ],
    )

    add_heading(doc, "Recommended Board View", 1)
    add_body(
        doc,
        "At the present stage, the application should be understood as a credible property operations platform that already supports meaningful customer workflows, but is still one roadmap phase away from being a fully automated digital housing transaction system."
    )
    add_body(
        doc,
        "The practical recommendation is to treat the current build as a strong operating foundation. The next investment should focus on real payment integration, deeper automation, lease execution, and buyer-finance tooling so that the platform can move from helpful workflow support to complete transaction completion."
    )

    doc.save(output_path)


def main():
    output_path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.getcwd(), "docs", f"sandworth-homes-board-brief-{date(2026, 8, 28).isoformat()}.docx"
    )
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    build_document(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
